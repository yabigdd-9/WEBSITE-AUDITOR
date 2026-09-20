#!/usr/bin/env python3
"""Export a Website Auditor portfolio to CSV, XLSX and PDF."""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


COLUMNS = [
    "domain",
    "url",
    "opportunity_score",
    "health_score",
    "finding_count",
    "profile",
    "site_type",
    "top_findings",
    "timestamp",
]


def load_audits(audits_dir: str | Path) -> list[dict[str, Any]]:
    root = Path(audits_dir)
    audits: list[dict[str, Any]] = []
    for path in sorted(root.glob("*.json")):
        try:
            payload = json.loads(path.read_text())
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(payload, dict) and payload.get("domain"):
            audits.append(payload)
    return audits


def flatten_audit(audit: dict[str, Any]) -> dict[str, Any]:
    findings = audit.get("site_findings") or audit.get("findings") or []
    labels = []
    for item in findings[:5] if isinstance(findings, list) else []:
        if isinstance(item, dict):
            labels.append(str(item.get("check_id") or item.get("message") or item.get("defect") or ""))

    category_scores = audit.get("site_category_scores") or audit.get("category_scores") or {}
    health = category_scores.get("overall_health_score") if isinstance(category_scores, dict) else None
    profile = (audit.get("audit_profile") or {}).get("name")
    site_type = (audit.get("site_type") or {}).get("site_type")

    return {
        "domain": audit.get("domain"),
        "url": audit.get("url"),
        "opportunity_score": audit.get("score"),
        "health_score": health,
        "finding_count": len(findings) if isinstance(findings, list) else audit.get("defect_count", 0),
        "profile": profile,
        "site_type": site_type,
        "top_findings": "; ".join(labels),
        "timestamp": audit.get("timestamp"),
    }


def portfolio_rows(audits: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = [flatten_audit(item) for item in audits]
    rows.sort(
        key=lambda row: (
            -(float(row["opportunity_score"]) if isinstance(row["opportunity_score"], (int, float)) else -1),
            str(row.get("domain") or ""),
        )
    )
    return rows


def export_csv(rows: list[dict[str, Any]], path: str | Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    return destination


def export_xlsx(rows: list[dict[str, Any]], path: str | Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    ws.append(COLUMNS)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([row.get(column) for column in COLUMNS])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "domain": 28,
        "url": 42,
        "opportunity_score": 18,
        "health_score": 14,
        "finding_count": 14,
        "profile": 18,
        "site_type": 22,
        "top_findings": 70,
        "timestamp": 28,
    }
    for index, column in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(index)].width = widths[column]

    wb.save(destination)
    return destination


def export_pdf(rows: list[dict[str, Any]], path: str | Path, *, title: str = "Website Audit Portfolio") -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(destination),
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
        title=title,
    )
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    table_data = [[
        "Domain",
        "Opportunity",
        "Health",
        "Findings",
        "Profile",
        "Site type",
        "Top findings",
    ]]
    for row in rows:
        table_data.append([
            str(row.get("domain") or ""),
            str(row.get("opportunity_score") if row.get("opportunity_score") is not None else "—"),
            str(row.get("health_score") if row.get("health_score") is not None else "—"),
            str(row.get("finding_count") or 0),
            str(row.get("profile") or ""),
            str(row.get("site_type") or ""),
            str(row.get("top_findings") or "")[:180],
        ])

    table = Table(table_data, repeatRows=1, colWidths=[90, 60, 50, 50, 70, 80, 310])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
    ]))
    story.append(table)
    doc.build(story)
    return destination


def export_all(audits_dir: str | Path, output_dir: str | Path, *, title: str) -> dict[str, str]:
    rows = portfolio_rows(load_audits(audits_dir))
    output = Path(output_dir)
    return {
        "csv": str(export_csv(rows, output / "portfolio.csv")),
        "xlsx": str(export_xlsx(rows, output / "portfolio.xlsx")),
        "pdf": str(export_pdf(rows, output / "portfolio.pdf", title=title)),
        "row_count": str(len(rows)),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Export Website Auditor portfolio reports")
    parser.add_argument("--audits-dir", default="audits")
    parser.add_argument("--output-dir", default="outputs/portfolio")
    parser.add_argument("--title", default="Website Audit Portfolio")
    parser.add_argument("--format", choices=["all", "csv", "xlsx", "pdf"], default="all")
    args = parser.parse_args()

    rows = portfolio_rows(load_audits(args.audits_dir))
    output = Path(args.output_dir)
    results: dict[str, str] = {}
    if args.format in {"all", "csv"}:
        results["csv"] = str(export_csv(rows, output / "portfolio.csv"))
    if args.format in {"all", "xlsx"}:
        results["xlsx"] = str(export_xlsx(rows, output / "portfolio.xlsx"))
    if args.format in {"all", "pdf"}:
        results["pdf"] = str(export_pdf(rows, output / "portfolio.pdf", title=args.title))
    results["row_count"] = str(len(rows))
    print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
